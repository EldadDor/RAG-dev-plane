from __future__ import annotations

import base64
import io
from pathlib import Path
from unittest.mock import AsyncMock
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
import pytest

from app.chunkers.chunker_adapter import DefaultChunker
from app.chunkers.powerpoint_chunker import chunk_powerpoint_document
from app.config import Settings
from app.domain.models import SourceType
from app.loaders.powerpoint_loader import PowerPointDocumentError, PowerPointLoader
from app.loaders.registry import load_document, load_directory, UnsupportedFileTypeError
from app.services.document_catalog import safe_display_metadata
from app.services.ingestion_service import IngestionService

PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=")


def write_deck(path, *, objects=False):
    deck = Presentation()
    deck.core_properties.title = "Operations presentation"
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Deployment"
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box.text = "Deploy safely.\nמדריך תפעול"
    box.text_frame.paragraphs[1].level = 1
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for cell, text in zip((table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)), ("Setting", "Value", "Port", "8000")):
        cell.text = text
    chart_data = CategoryChartData()
    chart_data.categories = ["East", "West"]
    chart_data.add_series("Requests", [42, 120])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(2), chart_data)
    picture = slide.shapes.add_picture(io.BytesIO(PNG), Inches(6), Inches(1), width=Inches(1))
    picture.element.xpath(".//p:cNvPr")[0].set("descr", "Readiness screenshot")
    slide.notes_slide.notes_text_frame.text = "Verify health before deployment."
    if objects:
        slide.shapes.add_ole_object(io.BytesIO(b"opaque binary"), "Excel.Sheet.12", Inches(6), Inches(3),
                                    width=Inches(1), height=Inches(1), icon_file=io.BytesIO(PNG))
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Rollback"
    slide.element.set("show", "0")
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Restore the previous release."
    slide.shapes.add_picture(io.BytesIO(PNG), Inches(6), Inches(1), width=Inches(1))
    deck.save(path)


def rewrite_package(path, replacements=None, additions=None):
    replacements, additions = replacements or {}, additions or {}
    with zipfile.ZipFile(path) as archive:
        entries = [(entry, archive.read(entry.filename)) for entry in archive.infolist()]
    with zipfile.ZipFile(path, "w") as archive:
        for entry, content in entries:
            change = replacements.get(entry.filename)
            archive.writestr(entry, change(content) if callable(change) else change if change is not None else content)
        for name, content in additions.items():
            archive.writestr(name, content)


def test_mixed_deck_structure_assets_notes_unicode_and_hidden_slides(tmp_path):
    source = tmp_path / "guide.pptx"
    write_deck(source)
    doc = load_document(str(source))
    assert doc.source_type == SourceType.powerpoint and doc.title == "Operations presentation"
    assert doc.metadata["total_slides"] == 2
    assert doc.content.index("Deploy safely.") < doc.content.index("| Setting | Value |") < doc.content.index("Series: Requests")
    assert "East: 42" in doc.content and "West: 120" in doc.content
    assert "מדריך תפעול" in doc.content and "Speaker notes:\nVerify health" in doc.content
    assert doc.metadata["slides"][1]["hidden"] is True
    assert len(doc.assets) == 2 and doc.assets[0].alt_text == "Readiness screenshot"
    assert doc.assets[0].content == PNG and doc.assets[0].media_type == "image/png"
    assert doc.assets[0].relationship_id != doc.assets[1].relationship_id
    assert doc.content[doc.assets[0].source_index:].startswith("[Image:")
    assert len(doc.content_hash) == 64
    assert load_directory(str(tmp_path))[1] == []
    assert safe_display_metadata({"source_path": str(source), "source_type": "powerpoint"})["document_type"] == "powerpoint"


def test_slide_chunking_keeps_page_offsets_and_assets_local(tmp_path):
    source = tmp_path / "guide.pptx"
    write_deck(source)
    doc = load_document(str(source))
    chunks = chunk_powerpoint_document(doc, DefaultChunker(chunk_size=90, chunk_overlap=12))
    assert {c.metadata["page"] for c in chunks} == {1, 2}
    for chunk in chunks:
        slide = doc.metadata["slides"][chunk.metadata["page"] - 1]
        assert slide["start_index"] <= chunk.start_index < chunk.end_index <= slide["end_index"]
        assert doc.content[chunk.start_index:chunk.end_index] == chunk.text
        assert chunk.metadata["section"] == slide["title"]
        assert chunk.metadata["hidden"] == slide["hidden"]


def test_xy_chart_preserves_x_and_y_cache(tmp_path):
    path = tmp_path / "scatter.pptx"
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = XyChartData(); series = data.add_series("Measured")
    series.add_data_point(2.5, 8.0); series.add_data_point(5.0, 13.5)
    slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(1), Inches(1), Inches(4), Inches(3), data)
    deck.save(path)
    text = load_document(str(path)).content
    assert "x=2.5, y=8.0" in text and "x=5.0, y=13.5" in text


def test_smartart_retains_cached_text_and_records_structure_limit(tmp_path):
    source = tmp_path / "smartart.pptx"
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    data = b'''<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dgm:ptLst><dgm:pt modelId="one"><dgm:t><a:p><a:r><a:t>Plan then deploy</a:t></a:r></a:p></dgm:t></dgm:pt></dgm:ptLst></dgm:dataModel>'''
    part = Part(PackURI('/ppt/diagrams/data1.xml'), 'application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml', slide.part.package, data)
    rid = slide.part.relate_to(part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData')
    shape = parse_xml(f'''<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:nvGraphicFramePr><p:cNvPr id="200" name="Process"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr><p:xfrm><a:off x="0" y="0"/><a:ext cx="1000" cy="1000"/></p:xfrm><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram"><dgm:relIds r:dm="{rid}"/></a:graphicData></a:graphic></p:graphicFrame>''')
    slide.shapes._spTree.append(shape); deck.save(source)
    doc = load_document(str(source))
    assert "Plan then deploy" in doc.content
    assert doc.metadata["extraction_warnings"][0]["code"] == "smartart_text_only"


def test_chart_without_cache_is_reported_without_inventing_values(tmp_path):
    source = tmp_path / "no-cache.pptx"; write_deck(source)
    def remove_cache(content):
        root = etree.fromstring(content)
        for element in list(root.iter()):
            if etree.QName(element).localname in {'numCache', 'strCache'}:
                element.getparent().remove(element)
        return etree.tostring(root)
    rewrite_package(source, {'ppt/charts/chart1.xml': remove_cache})
    doc = load_document(str(source))
    assert "[Chart cached data unavailable]" in doc.content and 'West: 120' not in doc.content
    assert any(w['code'] == 'chart_cache_unavailable' for w in doc.metadata['extraction_warnings'])


def test_ole_is_reported_without_opening_binary_and_preview_is_asset(tmp_path):
    source = tmp_path / "objects.pptx"
    write_deck(source, objects=True)
    doc = load_document(str(source))
    assert len(doc.metadata["unsupported_objects"]) == 1
    assert "content unavailable" in doc.content
    assert len(doc.assets) == 3
    assert all(asset.content == PNG for asset in doc.assets)


def test_image_only_change_updates_package_and_asset_hash(tmp_path):
    source = tmp_path / "guide.pptx"
    write_deck(source)
    first = load_document(str(source))
    with zipfile.ZipFile(source) as archive:
        media = next(n for n in archive.namelist() if n.startswith("ppt/media/"))
    rewrite_package(source, {media: PNG + b"changed"})
    changed = load_document(str(source))
    assert first.content == changed.content and first.content_hash != changed.content_hash
    assert first.assets[0].content_hash != changed.assets[0].content_hash


def test_blank_deck_is_empty_and_image_only_slide_is_indexable(tmp_path):
    source = tmp_path / "empty.pptx"
    deck = Presentation(); deck.slides.add_slide(deck.slide_layouts[6]); deck.save(source)
    doc = load_document(str(source))
    assert doc.content == "" and chunk_powerpoint_document(doc, DefaultChunker()) == []
    deck.slides[0].shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(1)); deck.save(source)
    doc = load_document(str(source))
    assert "[Image:" in doc.content and len(doc.assets) == 1
    assert chunk_powerpoint_document(doc, DefaultChunker())[0].metadata["page"] == 1


def test_external_picture_is_not_fetched(tmp_path):
    source = tmp_path / "external.pptx"
    write_deck(source)
    def externalize(content):
        root = etree.fromstring(content)
        for relation in root:
            if relation.get("Type", "").endswith("/image"):
                relation.set("TargetMode", "External")
                relation.set("Target", "https://unreachable.invalid/image.png")
        return etree.tostring(root)
    rewrite_package(source, {"ppt/slides/_rels/slide1.xml.rels": externalize})
    doc = load_document(str(source))
    assert len(doc.assets) == 1
    assert any(w["code"] == "image_unavailable_or_external" for w in doc.metadata["extraction_warnings"])


def test_external_chart_is_not_fetched_and_other_slide_content_survives(tmp_path):
    source = tmp_path / "external-chart.pptx"; write_deck(source)
    def externalize(content):
        root = etree.fromstring(content)
        for relation in root:
            if relation.get("Type", "").endswith("/chart"):
                relation.set("TargetMode", "External")
                relation.set("Target", "https://unreachable.invalid/chart.xml")
        return etree.tostring(root)
    rewrite_package(source, {"ppt/slides/_rels/slide1.xml.rels": externalize})
    doc = load_document(str(source))
    assert "[Chart cached data unavailable]" in doc.content and "Deploy safely." in doc.content
    assert any(w["code"] == "chart_cache_unavailable" for w in doc.metadata["extraction_warnings"])


@pytest.mark.parametrize("content, message", [(b"legacy or encrypted", "not a readable"), (b"PK", "not a readable")])
def test_invalid_packages_fail_without_publication(tmp_path, content, message):
    source = tmp_path / "bad.pptx"; source.write_bytes(content)
    with pytest.raises(PowerPointDocumentError, match=message):
        load_document(str(source))


def test_unsafe_and_oversized_packages_rejected(tmp_path, monkeypatch):
    source = tmp_path / "unsafe.pptx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("ppt/presentation.xml", "<presentation/>")
        archive.writestr("..\\escape.bin", b"unsafe")
    with pytest.raises(PowerPointDocumentError, match="unsafe entry"):
        load_document(str(source))
    write_deck(source)
    monkeypatch.setattr("app.loaders.powerpoint_loader._MAX_EXPANDED_BYTES", 10)
    with pytest.raises(PowerPointDocumentError, match="expanded-byte"):
        load_document(str(source))


def test_legacy_ppt_stays_unsupported_and_directory_reports_bad_pptx(tmp_path):
    legacy = tmp_path / "legacy.ppt"; legacy.write_bytes(b"old")
    with pytest.raises(UnsupportedFileTypeError):
        load_document(str(legacy))
    broken = tmp_path / "broken.pptx"; broken.write_bytes(b"bad")
    docs, skipped = load_directory(str(tmp_path))
    assert docs == [] and skipped[0]["path"] == str(broken)


@pytest.mark.asyncio
async def test_pptx_uses_normal_ingestion_publication_assets_and_dry_run(tmp_path):
    source = tmp_path / "guide.pptx"; write_deck(source)
    settings = Settings(_env_file=None, VECTOR_STORE="qdrant", CHUNKER_PROVIDER="default", CHUNK_SIZE=100, CHUNK_OVERLAP=12)
    vector_store, provider = AsyncMock(), AsyncMock()
    vector_store.get_document_hash.return_value = None
    provider.create_embedding.return_value = [0.1, 0.2, 0.3]
    service = IngestionService(settings, provider, vector_store)
    await service.ingest_path(str(source), dry_run=True)
    provider.create_embedding.assert_not_awaited(); vector_store.replace_document.assert_not_awaited()
    result = await service.ingest_path(str(source))
    document, chunks, assets = vector_store.replace_document.await_args.args
    assert document["source_type"] == "powerpoint" and result.documents[0].assets_found == 2
    assert {c["payload"]["page"] for c in chunks} == {1, 2}
    for asset in assets:
        linked = next(c for c in chunks if c["chunk_id"] in asset["related_chunk_ids"])
        assert linked["payload"]["section"] in asset["caption"]
    vector_store.get_document_hash.return_value = document["content_hash"]
    provider.reset_mock(); vector_store.replace_document.reset_mock()
    assert (await service.ingest_path(str(source))).documents[0].skip_reason == "unchanged"
    provider.create_embedding.assert_not_awaited(); vector_store.replace_document.assert_not_awaited()
