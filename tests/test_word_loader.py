from __future__ import annotations

import base64
from pathlib import Path
import zipfile

import pytest
from docx import Document as OpenWordDocument
from docx.enum.text import WD_BREAK
from docx.shared import Inches

from app.domain.models import SourceType
from app.loaders.registry import load_directory, load_document
from app.loaders.word_loader import WordDocumentError, WordLoader


_ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _write_word_fixture(path: Path, image_path: Path | None = None) -> None:
    document = OpenWordDocument()
    document.core_properties.title = "Operations Guide"
    document.add_heading("Deployment", level=1)
    document.add_paragraph("Deploy the service from the approved package.")
    document.add_paragraph("Verify readiness", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Setting"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Port"
    table.cell(1, 1).text = "8000"
    image_paragraph = document.add_paragraph("Console screenshot")
    if image_path is not None:
        image_paragraph.add_run().add_picture(str(image_path), width=Inches(0.25))
    page_break = document.add_paragraph()
    page_break.add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("Rollback", level=1)
    document.add_paragraph("Restore the previous package.")
    document.save(path)


def _replace_first_media_entry(path: Path, replacement: bytes) -> None:
    with zipfile.ZipFile(path) as package:
        entries = [(entry, package.read(entry.filename)) for entry in package.infolist()]
    with zipfile.ZipFile(path, "w") as package:
        replaced = False
        for entry, content in entries:
            if not replaced and entry.filename.startswith("word/media/"):
                content = replacement
                replaced = True
            package.writestr(entry, content)
    assert replaced


def test_word_loader_extracts_ordered_structure_and_image_anchor(tmp_path):
    image = tmp_path / "screen.png"
    image.write_bytes(_ONE_PIXEL_PNG)
    source = tmp_path / "guide.docx"
    _write_word_fixture(source, image)

    loaded = WordLoader().load(str(source))

    assert loaded.source_type == SourceType.word
    assert loaded.title == "Operations Guide"
    assert loaded.content.index("# Deployment") < loaded.content.index("Deploy the service")
    assert loaded.content.index("Deploy the service") < loaded.content.index("| Setting | Value |")
    assert "- Verify readiness" in loaded.content
    assert "[PAGE BREAK]" in loaded.content
    assert loaded.content.index("[PAGE BREAK]") < loaded.content.index("# Rollback")
    assert loaded.metadata["image_count"] == 1
    anchor = loaded.metadata["image_anchors"][0]
    assert anchor["available"] is True
    assert anchor["media_type"] == "image/png"
    assert anchor["byte_size"] == len(_ONE_PIXEL_PNG)
    assert anchor["section"] == "Deployment"
    assert anchor["caption"] == "Console screenshot"
    assert anchor["display_name"] == "Picture 1"
    assert anchor["block_id"].startswith("block-")
    assert len(loaded.content_hash or "") == 64


def test_word_loader_hash_changes_when_only_image_bytes_change(tmp_path):
    image = tmp_path / "screen.png"
    image.write_bytes(_ONE_PIXEL_PNG)
    source = tmp_path / "guide.docx"
    _write_word_fixture(source, image)
    original = WordLoader().load(str(source))

    _replace_first_media_entry(source, b"replacement-image-bytes")
    changed = WordLoader().load(str(source))

    assert changed.content == original.content
    assert changed.content_hash != original.content_hash
    assert changed.metadata["image_anchors"][0]["content_hash"] != original.metadata["image_anchors"][0]["content_hash"]


def test_registry_and_directory_discovery_support_docx(tmp_path):
    source = tmp_path / "guide.docx"
    _write_word_fixture(source)

    loaded = load_document(str(source))
    documents, skipped = load_directory(str(tmp_path))

    assert loaded.source_type == SourceType.word
    assert [document.source_path for document in documents] == [str(source)]
    assert skipped == []


def test_word_loader_rejects_renamed_non_docx(tmp_path):
    source = tmp_path / "fake.docx"
    source.write_bytes(b"not a Word package")

    with pytest.raises(WordDocumentError, match="not a readable .docx"):
        WordLoader().load(str(source))


def test_word_loader_rejects_package_entry_traversal(tmp_path):
    source = tmp_path / "unsafe.docx"
    with zipfile.ZipFile(source, "w") as package:
        package.writestr("[Content_Types].xml", "<Types />")
        package.writestr("word/document.xml", "<document />")
        package.writestr("../escaped.bin", b"unsafe")

    with pytest.raises(WordDocumentError, match="unsafe entry path"):
        WordLoader().load(str(source))


def test_word_loader_rejects_zip_without_word_parts(tmp_path):
    source = tmp_path / "renamed.docx"
    with zipfile.ZipFile(source, "w") as package:
        package.writestr("notes.txt", "not a Word document")

    with pytest.raises(WordDocumentError, match="not a valid .docx"):
        WordLoader().load(str(source))
