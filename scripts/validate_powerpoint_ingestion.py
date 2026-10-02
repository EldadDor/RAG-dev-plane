"""Provider-free local PowerPoint smoke check; all SQL fixture writes roll back."""
from __future__ import annotations

import argparse
import asyncio
import base64
import io
import json
from pathlib import Path
import sys
import tempfile
import uuid
from datetime import datetime, timezone

import httpx
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from app.config import get_settings
from app.main import _init_pg_vector_store
from app.services.asset_store import InMemoryAssetStore
from app.services.document_catalog import DocumentCatalogService, DocumentCursor, PinnedConnectionPool, PostgresDocumentCatalog
from app.services.ingestion_service import IngestionService
from app.services.model_profiles import PostgresModelProfileStore, InMemoryEmbeddingCache

PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=")


def write_fixture(path, version="one", blank=False):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    if not blank:
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "מדריך תפעול " + version
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(3), Inches(1)).table
        for row, values in enumerate((("Region", "Sales"), ("East", "42"))):
            for column, value in enumerate(values):
                table.cell(row, column).text = value
        data = CategoryChartData()
        data.categories = ["East", "West"]
        data.add_series("Revenue", [42, 120])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(2), data)
        slide.shapes.add_picture(io.BytesIO(PNG), Inches(6), Inches(1))
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Rollback " + version
        slide.notes_slide.notes_text_frame.text = "Use the previous release."
    deck.save(path)


async def validate():
    settings = get_settings()
    assert settings.app_env == "local" and settings.vector_store == "postgres"
    assert settings.chunking_profile(None)[1].provider in {
        "default", "langchain", "recursive-character", "chonkie-recursive", "recursive"
    }, "Provider-free validation requires a non-semantic chunking profile"
    report = {"validated_at": datetime.now(timezone.utc).isoformat(), "real_provider_calls": 0, "checks": []}
    store = await _init_pg_vector_store(settings)
    schema = settings.pg_schema
    try:
        async with store.pool.acquire() as conn:
            transaction = conn.transaction(isolation="repeatable_read")
            await transaction.start()
            workspace = "np21-qa-" + uuid.uuid4().hex
            try:
                await conn.execute(f"INSERT INTO {schema}.workspaces(workspace_id,display_name) VALUES($1,'PowerPoint QA')", workspace)
                pinned = PinnedConnectionPool(conn)
                profiles = PostgresModelProfileStore(pinned, schema)
                profile = await profiles.get(settings.model_profile)

                class SyntheticEmbedding:
                    fail = False

                    async def create_embedding(self, model, text):
                        if self.fail:
                            raise RuntimeError("Intentional QA provider failure")
                        return [0.01] * profile.dimensions

                provider = SyntheticEmbedding()
                assets = InMemoryAssetStore()
                service = IngestionService(settings, provider, store.bound_to(conn), assets, profiles, InMemoryEmbeddingCache())
                with tempfile.TemporaryDirectory(prefix="np21-pptx-") as temp:
                    source = Path(temp) / "fixture.pptx"
                    write_fixture(source)
                    result = await service.ingest_path(str(source), workspace_id=workspace)
                    doc_id = result.documents[0].doc_id
                    assert result.chunks_indexed >= 2 and result.documents[0].assets_found == 1
                    count = await conn.fetchval(f"SELECT count(*) FROM {schema}.{profile.storage_target} WHERE metadata->>'workspace_id'=$1", workspace)
                    assert count == result.chunks_indexed
                    rows = await conn.fetch(f"SELECT metadata FROM {schema}.{profile.storage_target} WHERE metadata->>'workspace_id'=$1", workspace)
                    assert {r["metadata"]["page"] for r in rows} == {1, 2}
                    assert all(r["metadata"]["source_type"] == "powerpoint" for r in rows)
                    catalog = DocumentCatalogService(settings, PostgresDocumentCatalog(pinned, schema), profiles, DocumentCursor("np21-qa-key-" * 4))
                    page = await catalog.list(workspace, "qa")
                    item = page.items[0]
                    assert item.document_type == "powerpoint" and item.indexed_chunk_count == count
                    assert await conn.fetchval(f"SELECT count(*) FROM {schema}.document_index_assets WHERE workspace_id=$1", workspace) == 1
                    assert (await service.ingest_path(str(source), workspace_id=workspace)).documents[0].skip_reason == "unchanged"
                    stamp = item.last_ingested_at
                    assert (await catalog.list(workspace, "qa")).items[0].last_ingested_at == stamp
                    report["checks"].extend(["mixed_deck_real_sql_publication", "slide_citation_numbers", "catalog_powerpoint_type_count",
                                             "image_asset_ownership", "unchanged_skip_preserves_time"])

                    write_fixture(source, version="changed")
                    provider.fail = True
                    try:
                        await service.ingest_path(str(source), workspace_id=workspace)
                        raise AssertionError("Expected provider failure")
                    except RuntimeError as exc:
                        assert str(exc) == "Intentional QA provider failure"
                    assert (await catalog.list(workspace, "qa")).items[0].last_ingested_at == stamp
                    provider.fail = False
                    await service.ingest_path(str(source), workspace_id=workspace)
                    assert (await catalog.list(workspace, "qa")).items[0].last_ingested_at > stamp
                    write_fixture(source, blank=True)
                    await service.ingest_path(str(source), workspace_id=workspace)
                    assert (await catalog.list(workspace, "qa")).items[0].indexed_chunk_count == 0
                    assert await conn.fetchval(f"SELECT count(*) FROM {schema}.document_assets WHERE workspace_id=$1", workspace) == 0
                    report["checks"].extend(["failed_reingestion_preserves_publication", "changed_reingestion", "blank_deck_removes_chunks_assets"])
                    report["initial_chunks"] = count
            finally:
                await transaction.rollback()
            assert not await conn.fetchval(f"SELECT EXISTS(SELECT 1 FROM {schema}.workspaces WHERE workspace_id=$1)", workspace)
            report["fixtures_rolled_back"] = True
        async with httpx.AsyncClient(base_url="http://127.0.0.1:8000", timeout=15) as client:
            assert (await client.get("/health")).status_code == 200
            assert (await client.get("/readiness")).status_code == 200
            schema_response = (await client.get("/openapi.json")).json()
            assert "powerpoint" in schema_response["components"]["schemas"]["DocumentSummary"]["properties"]["document_type"]["enum"]
            with tempfile.TemporaryDirectory(prefix="np21-http-pptx-") as temp:
                source = Path(temp) / "fixture.pptx"
                write_fixture(source)
                response = await client.post("/ingest", json={"source_path": str(source), "dry_run": True})
                assert response.status_code == 200, response.text
                result = response.json()
                assert result["dry_run"] is True and result["indexed"] >= 2
                assert result["documents"][0]["assets_found"] == 1
        report["checks"].append("updated_http_health_readiness_schema")
        report["checks"].append("live_http_pptx_ingest_dry_run")
        return report
    finally:
        await store.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    encoded = json.dumps(asyncio.run(validate()), indent=2, sort_keys=True)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(encoded + "\n", encoding="utf-8")
    print(encoded)
