"""Extract slide-local text/structure and image assets without rendering slides."""
from __future__ import annotations

import hashlib
from pathlib import Path, PurePosixPath
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.chunkers.ids import make_doc_id
from app.domain.models import Document, DocumentAsset, SourceType

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_DGM = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}"
_MAX_PACKAGE_ENTRIES = 10_000
_MAX_EXPANDED_BYTES = 250 * 1024 * 1024


class PowerPointDocumentError(ValueError):
    """The input is not a safe, readable modern PowerPoint presentation."""


def _validate_package(path: Path):
    try:
        with zipfile.ZipFile(path) as package:
            entries = package.infolist()
            if len(entries) > _MAX_PACKAGE_ENTRIES:
                raise PowerPointDocumentError("PowerPoint package has too many entries")
            if sum(entry.file_size for entry in entries) > _MAX_EXPANDED_BYTES:
                raise PowerPointDocumentError("PowerPoint package exceeds expanded-byte safety limit")
            names = [entry.filename for entry in entries]
            if len(names) != len(set(names)):
                raise PowerPointDocumentError("PowerPoint package contains duplicate entry paths")
            for name in names:
                normalized = name.replace("\\", "/")
                if PurePosixPath(normalized).is_absolute() or ".." in PurePosixPath(normalized).parts or ":" in normalized:
                    raise PowerPointDocumentError("PowerPoint package contains an unsafe entry path")
            if not {"[Content_Types].xml", "ppt/presentation.xml"} <= set(names):
                raise PowerPointDocumentError("File is not a valid .pptx package")
            if any(entry.flag_bits & 1 for entry in entries):
                raise PowerPointDocumentError("Encrypted PowerPoint packages are unsupported")
            if package.testzip() is not None:
                raise PowerPointDocumentError("PowerPoint package contains a corrupt entry")
    except (zipfile.BadZipFile, RuntimeError, NotImplementedError) as exc:
        raise PowerPointDocumentError("File is not a readable .pptx package; encrypted and legacy .ppt files are unsupported") from exc


def _text(element):
    return " ".join(node.text or "" for node in element.iter(_A + "t")).strip()


def _table_text(table):
    rows = [["" if cell.is_spanned else cell.text.strip().replace("\n", " ").replace("|", "\\|")
             for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    rendered = ["| " + " | ".join(row) + " |" for row in rows]
    rendered.insert(1, "| " + " | ".join("---" for _ in rows[0]) + " |")
    return "\n".join(rendered)


def _cached_values(parent):
    """Preserve point indices, including gaps in category/XY/bubble caches."""
    if parent is None:
        return {}
    result = {}
    for point in parent.iter(_C + "pt"):
        value = point.find(_C + "v")
        if value is not None:
            index = int(point.get("idx", "0"))
            text = value.text or ""
            result[index] = result[index] + " / " + text if index in result else text
    return result


def _chart_text(chart):
    # Read OOXML caches, never linked workbooks. This also works for XY/bubble
    # series and mixed plots without relying on category-only chart APIs.
    root = chart._chartSpace
    title = root.find(".//" + _C + "title")
    title_text = _text(title) if title is not None else ""
    if title is not None and not title_text:
        title_text = " ".join(v.text or "" for v in title.iter(_C + "v"))
    lines = ["Chart: " + (title_text or "Untitled chart")]
    points_found = 0
    for ordinal, series in enumerate(root.iter(_C + "ser"), 1):
        tx = series.find(_C + "tx")
        name = " ".join(v.text or "" for v in tx.iter(_C + "v")) if tx is not None else ""
        lines.append("Series: " + (name or f"{ordinal}"))
        categories = _cached_values(series.find(_C + "cat"))
        xs = _cached_values(series.find(_C + "xVal"))
        values = _cached_values(series.find(_C + "val")) or _cached_values(series.find(_C + "yVal"))
        sizes = _cached_values(series.find(_C + "bubbleSize"))
        for index in sorted(set(categories) | set(xs) | set(values) | set(sizes)):
            label = categories.get(index, f"Point {index + 1}")
            if xs:
                line = f"{label}: x={xs.get(index, 'unavailable')}, y={values.get(index, 'unavailable')}"
            else:
                line = f"{label}: {values.get(index, 'unavailable')}"
            if sizes:
                line += f", size={sizes.get(index, 'unavailable')}"
            lines.append(line)
            points_found += 1
    if not points_found:
        lines.append("[Chart cached data unavailable]")
    return "\n".join(lines), bool(points_found)


def _diagram_text(shape, slide):
    texts = []
    for refs in shape.element.iter(_DGM + "relIds"):
        relation = slide.part.rels.get(refs.get(_R + "dm"))
        if relation is None or relation.is_external:
            continue
        root = etree.fromstring(relation.target_part.blob,
                                etree.XMLParser(resolve_entities=False, no_network=True))
        for point in root.iter(_DGM + "pt"):
            text = _text(point)
            if text:
                texts.append(text)
    return "\n".join(texts)


class PowerPointLoader:
    """Load slide order, text, tables, cached charts, notes and embedded images."""

    def load(self, source_path: str) -> Document:
        path = Path(source_path)
        _validate_package(path)
        try:
            presentation = Presentation(str(path))
            return self._extract(presentation, path, source_path)
        except PowerPointDocumentError:
            raise
        except Exception as exc:
            raise PowerPointDocumentError("Unable to parse PowerPoint presentation") from exc

    def _extract(self, presentation, path, source_path):
        parts, blocks, slides, assets, anchors, warnings, unsupported = [], [], [], [], [], [], []
        offset = 0

        def append(text, kind, number, section, shape_id=None):
            nonlocal offset
            if parts:
                offset += 2
            block = {"block_id": f"block-{len(blocks):04d}", "ordinal": len(blocks),
                     "kind": kind, "slide_number": number, "section": section,
                     "shape_id": shape_id, "start_index": offset, "end_index": offset + len(text)}
            parts.append(text)
            blocks.append(block)
            offset += len(text)
            return block

        for number, slide in enumerate(presentation.slides, 1):
            title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
            section = title or f"Slide {number}"
            hidden = slide.element.get("show") in {"0", "false"}
            start_block = len(blocks)
            start = offset + (2 if parts else 0)

            def walk(shapes):
                for shape in shapes:
                    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                        yield from walk(shape.shapes)
                    else:
                        yield shape

            for shape in walk(slide.shapes):
                properties = next(shape.element.iter(_P + "cNvPr"), None)
                alt = (properties.get("descr") or properties.get("title")) if properties is not None else None
                kind, text = "drawing", ""
                if shape.has_text_frame:
                    kind = "title" if shape == slide.shapes.title else "text"
                    text = "\n".join(("  " * p.level) + p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                elif shape.has_table:
                    kind, text = "table", _table_text(shape.table)
                elif shape.has_chart:
                    kind = "chart"
                    try:
                        text, available = _chart_text(shape.chart)
                    except (ValueError, KeyError, NotImplementedError, AttributeError):
                        # External/unrecognized chart parts are not resolved or
                        # interpreted; retain the rest of the slide explicitly.
                        text, available = "[Chart cached data unavailable]", False
                    if not available:
                        warnings.append({"slide_number": number, "shape_id": shape.shape_id, "code": "chart_cache_unavailable"})
                elif any(True for _ in shape.element.iter(_DGM + "relIds")):
                    kind, text = "smartart", _diagram_text(shape, slide)
                    warnings.append({"slide_number": number, "shape_id": shape.shape_id, "code": "smartart_text_only"})
                    text = text or "[SmartArt text unavailable]"
                else:
                    text = _text(shape.element)
                    if (shape.shape_type in {MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT, MSO_SHAPE_TYPE.MEDIA}
                            or shape.element.tag == _P + "graphicFrame"):
                        kind = "unsupported_object"
                        unsupported.append({"slide_number": number, "shape_id": shape.shape_id,
                                            "shape_type": str(shape.shape_type), "name": shape.name})
                        text = text or f"[Embedded object: {shape.name}; content unavailable]"
                image_refs = list(shape.element.iter(_A + "blip"))
                if not text and image_refs:
                    kind, text = "image", f"[Image: {alt or shape.name}]"
                if not text:
                    continue
                if len(blocks) == start_block:
                    append(f"[SLIDE {number}] {section}", "slide", number, section)
                block = append(text, kind, number, section, shape.shape_id)
                for blip in image_refs:
                    rid = blip.get(_R + "embed") or blip.get(_R + "link")
                    relation = slide.part.rels.get(rid)
                    anchor_id = f"slide-{number:04d}-image-{len(anchors):04d}"
                    available = relation is not None and not relation.is_external and relation.reltype.endswith("/image")
                    anchor = {"anchor_id": anchor_id, "block_id": block["block_id"], "slide_number": number,
                              "relationship_id": rid, "available": available, "alt_text": alt}
                    anchors.append(anchor)
                    if not available:
                        warnings.append({"slide_number": number, "shape_id": shape.shape_id, "code": "image_unavailable_or_external"})
                        continue
                    part = relation.target_part
                    blob = part.blob
                    assets.append(DocumentAsset(
                        anchor_id=anchor_id, relationship_id=f"slide-{number}:{rid}", content=blob,
                        content_hash=hashlib.sha256(blob).hexdigest(), media_type=part.content_type,
                        original_name=PurePosixPath(str(part.partname)).name, ordinal=len(assets),
                        block_id=block["block_id"], block_ordinal=block["ordinal"], source_index=block["start_index"],
                        section=section, alt_text=alt, caption=f"Slide {number}: {section}"))
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                notes = frame.text.strip() if frame is not None else ""
                if notes:
                    if len(blocks) == start_block:
                        append(f"[SLIDE {number}] {section}", "slide", number, section)
                    append("Speaker notes:\n" + notes, "notes", number, section)
            slides.append({"slide_number": number, "title": section, "hidden": hidden,
                           "start_index": start if len(blocks) > start_block else offset, "end_index": offset})
        return Document(
            doc_id=make_doc_id(source_path), source_path=source_path, source_type=SourceType.powerpoint,
            content="\n\n".join(parts), title=(presentation.core_properties.title or "").strip() or path.stem,
            content_hash=hashlib.sha256(path.read_bytes()).hexdigest(), assets=assets,
            metadata={"document_format": "pptx", "total_slides": len(slides), "slides": slides,
                      "blocks": blocks, "image_anchors": anchors, "image_count": len(assets),
                      "unsupported_objects": unsupported, "extraction_warnings": warnings})
